"""Create professional Microsoft Word reports for static APK assessments."""

from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from analyzer.apk_analyzer import APKAnalysisResult

NAVY, BLUE, MUTED, TABLE_HEADER = "10243F", "2563EB", "64748B", "E8EEF5"
RISK_COLORS = {"CRITICAL": "B91C1C", "HIGH": "C2410C", "MEDIUM": "A16207", "LOW": "2563EB", "INFO": "475569"}


def _shade(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def _cell(cell, text: str, bold: bool = False, color: str = "152238", size: int = 10) -> None:
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(str(text))
    run.bold = bold
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def _table(document: Document, widths: list[float], headers: tuple[str, ...] | None = None):
    table = document.add_table(rows=1 if headers else 0, cols=len(widths))
    table.style = "Table Grid"
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for column, width in zip(table.columns, widths):
        column.width = Inches(width)
    if headers:
        for cell, text in zip(table.rows[0].cells, headers):
            _shade(cell, TABLE_HEADER)
            _cell(cell, text, bold=True, color=NAVY, size=9)
    return table


def _heading(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(14)
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(text)
    run.bold = True
    run.font.name = "Aptos Display"
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor.from_string(BLUE)


def _bullet(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.add_run(text)


def _metadata_table(document: Document, rows: list[tuple[str, str]]) -> None:
    table = _table(document, [1.7, 4.8])
    for label, value in rows:
        cells = table.add_row().cells
        _shade(cells[0], TABLE_HEADER)
        _cell(cells[0], label, bold=True, color=NAVY)
        _cell(cells[1], value)


def generate_report(result: APKAnalysisResult, sha256: str, output_folder: Path) -> Path:
    """Save the latest assessment as a formatted, readable .docx report."""
    output_folder.mkdir(parents=True, exist_ok=True)
    safe_name = result.metadata.filename.removesuffix(".apk").replace(" ", "_")
    path = output_folder / f"{safe_name}_static_assessment_{datetime.now():%Y%m%d_%H%M%S}.docx"
    metadata = result.metadata
    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.75)
    section.left_margin = section.right_margin = Inches(0.8)
    normal = document.styles["Normal"]
    normal.font.name, normal.font.size = "Aptos", Pt(10.5)
    normal.paragraph_format.space_after, normal.paragraph_format.line_spacing = Pt(6), 1.1

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_run = header.add_run("APK THREAT ANALYZER  |  STATIC ASSESSMENT")
    header_run.font.name, header_run.font.size = "Aptos", Pt(8)
    header_run.font.color.rgb = RGBColor.from_string(MUTED)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("Static assessment only - this APK was not installed or executed.")
    footer_run.font.name, footer_run.font.size = "Aptos", Pt(8)
    footer_run.font.color.rgb = RGBColor.from_string(MUTED)

    title = document.add_paragraph()
    title.paragraph_format.space_after = Pt(3)
    title_run = title.add_run("APK SECURITY ASSESSMENT REPORT")
    title_run.bold, title_run.font.name, title_run.font.size = True, "Aptos Display", Pt(24)
    title_run.font.color.rgb = RGBColor.from_string(NAVY)
    subtitle = document.add_paragraph(f"Static analysis of {metadata.filename} | Generated {datetime.now():%d %B %Y, %H:%M}")
    subtitle.paragraph_format.space_after = Pt(14)
    subtitle.runs[0].font.color.rgb = RGBColor.from_string(MUTED)

    _heading(document, "Executive Summary")
    paragraph = document.add_paragraph()
    risk = paragraph.add_run(f"{result.risk.level} RISK ({result.risk.score}/100). ")
    risk.bold, risk.font.color.rgb = True, RGBColor.from_string(RISK_COLORS.get(result.risk.level, MUTED))
    paragraph.add_run("This score is based on configured static indicators. It is not a malware verdict and must be reviewed in context.")
    assessment = _table(document, [1.45, 1.2, 3.85], ("Risk score", "Threat level", "Assessment basis"))
    cells = assessment.add_row().cells
    _cell(cells[0], f"{result.risk.score}/100", True, RISK_COLORS.get(result.risk.level, MUTED), 13)
    _cell(cells[1], result.risk.level, True, RISK_COLORS.get(result.risk.level, MUTED))
    _cell(cells[2], f"{len(result.findings)} configured findings; raw score {result.risk.raw_score} before the 100-point cap.")

    _heading(document, "APK Information")
    _metadata_table(document, [("File name", metadata.filename), ("Package name", metadata.package_name), ("Application", metadata.application_name), ("Version", f"{metadata.version_name} ({metadata.version_code})"), ("SHA-256", sha256), ("File size", f"{metadata.file_size_bytes:,} bytes"), ("Minimum SDK", metadata.min_sdk), ("Target SDK", metadata.target_sdk), ("Main activity", metadata.main_activity or "Not found")])

    _heading(document, "Risk Contributors")
    if result.risk.contributions:
        for finding in result.risk.contributions:
            _bullet(document, f"+{finding.score}: {finding.indicator} ({finding.category})")
    else:
        document.add_paragraph("No configured indicators contributed points to the score.")

    _heading(document, "Permissions Requested")
    if result.manifest.permissions:
        for permission in result.manifest.permissions:
            _bullet(document, permission)
    else:
        document.add_paragraph("No permissions were found in the manifest.")

    _heading(document, "Heuristic Threat Findings")
    findings_table = _table(document, [0.75, 1.1, 2.0, 0.55, 2.1], ("Severity", "Category", "Indicator", "Score", "Reason"))
    if result.findings:
        for finding in result.findings:
            cells = findings_table.add_row().cells
            _cell(cells[0], finding.severity, True, RISK_COLORS.get(finding.severity, MUTED), 9)
            _cell(cells[1], finding.category, size=9)
            _cell(cells[2], finding.indicator, size=9)
            _cell(cells[3], f"+{finding.score}", size=9)
            _cell(cells[4], finding.description, size=9)
    else:
        cells = findings_table.add_row().cells
        for cell, value in zip(cells, ("INFO", "None", "No configured indicators", "0", "No findings does not prove that the APK is safe.")):
            _cell(cell, value, size=9)

    _heading(document, "Android Components")
    _metadata_table(document, [("Activities", ", ".join(result.manifest.activities) or "None found"), ("Services", ", ".join(result.manifest.services) or "None found"), ("Receivers", ", ".join(result.manifest.receivers) or "None found"), ("Providers", ", ".join(result.manifest.providers) or "None found")])

    _heading(document, "Network Indicators")
    network = [item.indicator for item in result.findings if any(word in item.category.lower() for word in ("endpoint", "ip", "domain"))]
    if network:
        for item in network:
            _bullet(document, item)
    else:
        document.add_paragraph("No configured network indicators were found.")

    _heading(document, "Limitations and Safe Use")
    document.add_paragraph("This report describes static indicators collected from the APK file. It cannot prove that an APK is malicious or safe. The application did not install, launch, or execute the APK. Review the findings alongside the application's purpose and use additional security checks when required.")
    document.save(path)
    return path
